"""Resume bytes -> plain text. PDF (PyMuPDF, EasyOCR for image-only pages), DOCX, plain text.

Bad input raises ResumeParseError with a stable code; nothing here writes to disk.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field

from src.models.schemas import DocumentFormat, ErrorDetail, ErrorResponse

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_PAGES = 10
MIN_TEXT_CHARS = 30   # a PDF page with fewer non-space characters is treated as scanned
OCR_DPI = 200

_PDF_MAGIC = b"%PDF"
_ZIP_MAGIC = b"PK\x03\x04"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"  # legacy .doc and password-protected Office files


class ResumeParseError(Exception):
    """Structured, user-facing parse failure."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message

    def to_response(self) -> ErrorResponse:
        return ErrorResponse(error=ErrorDetail(code=self.code, message=self.message))


@dataclass
class ParsedDocument:
    text: str
    format: DocumentFormat
    pages: int = 1
    ocr_pages: list[int] = field(default_factory=list)


def parse_document(data: bytes, filename: str | None = None) -> ParsedDocument:
    """Detect the format from the bytes (filename only breaks ties) and extract text."""
    if not data:
        raise ResumeParseError("EMPTY_DOCUMENT", "The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ResumeParseError("FILE_TOO_LARGE", f"File is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    name = (filename or "").lower()
    if _PDF_MAGIC in data[:1024]:
        return _parse_pdf(data)
    if data.startswith(_ZIP_MAGIC) and (name.endswith(".docx") or _is_docx(data)):
        return _parse_docx(data)
    if data.startswith(_OLE_MAGIC):
        raise ResumeParseError("UNSUPPORTED_FORMAT",
                               "Legacy .doc or password-protected Office file. Save it as PDF or .docx and retry.")
    if name.endswith((".pdf", ".docx")):
        raise ResumeParseError("CORRUPT_FILE", "The file could not be read; it may be damaged.")
    if b"\x00" in data[:4096]:
        raise ResumeParseError("UNSUPPORTED_FORMAT", "Unsupported file type. Upload a PDF, .docx or .txt resume.")
    return parse_text(data.decode("utf-8", errors="replace"))


def parse_text(text: str) -> ParsedDocument:
    if not text or not text.strip():
        raise ResumeParseError("EMPTY_DOCUMENT", "No text found in the resume.")
    return ParsedDocument(text=text, format="text")


def _is_docx(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return "word/document.xml" in z.namelist()
    except zipfile.BadZipFile:
        return False


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise ResumeParseError("CORRUPT_FILE", "The .docx file could not be read; it may be damaged.") from None
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:  # many resume templates put sections in tables
        for row in table.rows:
            lines.append(" | ".join(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip())))
    text = "\n".join(lines)
    if not text.strip():
        raise ResumeParseError("EMPTY_DOCUMENT", "No text found in the .docx file.")
    return ParsedDocument(text=text, format="docx")


def _parse_pdf(data: bytes) -> ParsedDocument:
    import pymupdf as fitz

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise ResumeParseError("CORRUPT_FILE", "The PDF could not be read; it may be damaged.") from None
    with doc:
        if doc.needs_pass:
            raise ResumeParseError("ENCRYPTED_FILE", "The PDF is password-protected. Upload an unlocked copy.")
        if doc.page_count == 0:
            raise ResumeParseError("CORRUPT_FILE", "The PDF has no pages.")
        if doc.page_count > MAX_PAGES:
            raise ResumeParseError("TOO_MANY_PAGES", f"The PDF has more than {MAX_PAGES} pages.")
        texts, ocr_pages = [], []
        try:
            for number, page in enumerate(doc, start=1):
                text = page.get_text("text", sort=True)
                if len(re.sub(r"\s", "", text)) < MIN_TEXT_CHARS:
                    text = _ocr_page(page)
                    ocr_pages.append(number)
                texts.append(text)
        except ResumeParseError:
            raise
        except Exception:
            raise ResumeParseError("CORRUPT_FILE", "The PDF could not be read; it may be damaged.") from None
        pages = doc.page_count
    full = "\n".join(texts)
    if len(re.sub(r"\s", "", full)) < MIN_TEXT_CHARS:
        raise ResumeParseError("EMPTY_DOCUMENT", "No readable text found in the PDF, even with OCR.")
    return ParsedDocument(text=full, format="pdf", pages=pages, ocr_pages=ocr_pages)


_reader = None


def _ocr_reader():
    """EasyOCR reader, created on first use (loading it takes seconds)."""
    global _reader
    if _reader is None:
        import easyocr

        _reader = easyocr.Reader(["en"], gpu=False, verbose=False)
    return _reader


def _ocr_page(page) -> str:
    png = page.get_pixmap(dpi=OCR_DPI).tobytes("png")
    try:
        boxes = _ocr_reader().readtext(png, detail=1, paragraph=False)
    except Exception:
        raise ResumeParseError("OCR_FAILED", "A scanned page could not be read.") from None
    # Group text boxes into lines by vertical centre, then read each line left to right.
    rows: list[dict] = []
    for box, text, _conf in sorted(boxes, key=lambda b: (b[0][0][1] + b[0][2][1]) / 2):
        y = (box[0][1] + box[2][1]) / 2
        height = max(box[2][1] - box[0][1], 1)
        if rows and abs(rows[-1]["y"] - y) < height * 0.5:
            rows[-1]["items"].append((box[0][0], text))
        else:
            rows.append({"y": y, "items": [(box[0][0], text)]})
    return "\n".join(" ".join(t for _, t in sorted(r["items"])) for r in rows)
