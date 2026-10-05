"""Regenerate the test resumes: python module-2-skill-gap/tests/fixtures/make_fixtures.py

The person and companies are fictional placeholders.
"""
from pathlib import Path

import docx
import pymupdf as fitz

HERE = Path(__file__).resolve().parent

RESUME = """Aarav Mehta
Bengaluru | aarav.mehta@example.com | github.com/aarav-example
SUMMARY
Data analyst who will go the extra mile. Background in statistical analysis.
EXPERIENCE
Data Analyst | Swiggy | Jun 2021 - Present
- Built SQL dashboards in Tableau for city operations teams
- Automated weekly reports with Python and pandas
Junior Data Analyst, Mu Sigma | Jul 2019 - Aug 2021
- Wrote SQL queries and Excel models for retail clients
Data Science Intern | Fractal Analytics | Jan 2019 - Jun 2019
- Cleaned survey data in R
PROJECTS
Churn Prediction
- Built a churn model in Python with XGBoost and scikit-learn
- Served predictions through a FastAPI service
SKILLS
Python, SQL, Go, R, Excel, Docker, Power BI, Tableau
EDUCATION
B.Tech in Computer Science, NIT Trichy, 2019
"""


def text_pdf() -> fitz.Document:
    doc = fitz.open()
    page = doc.new_page()
    y = 50
    for line in RESUME.splitlines():
        page.insert_text((50, y), line, fontsize=11, fontname="helv")
        y += 17
    return doc


def main() -> None:
    doc = text_pdf()
    doc.save(HERE / "resume_text.pdf")

    # Scanned: the same page rendered to an image, with no text layer.
    pix = doc[0].get_pixmap(dpi=150)
    scanned = fitz.open()
    page = scanned.new_page(width=doc[0].rect.width, height=doc[0].rect.height)
    page.insert_image(page.rect, stream=pix.tobytes("jpg", jpg_quality=80))
    scanned.save(HERE / "resume_scanned.pdf")

    doc.save(HERE / "resume_encrypted.pdf", encryption=fitz.PDF_ENCRYPT_AES_256,
             user_pw="secret", owner_pw="owner-secret")

    data = (HERE / "resume_text.pdf").read_bytes()
    (HERE / "resume_corrupt.pdf").write_bytes(data[:200])  # truncated: header but no body

    d = docx.Document()
    for line in RESUME.splitlines():
        d.add_paragraph(line)
    d.save(HERE / "resume.docx")

    (HERE / "resume.txt").write_text(RESUME, encoding="utf-8")


if __name__ == "__main__":
    main()
